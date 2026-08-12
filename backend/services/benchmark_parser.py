"""
benchmark_parser.py
Parses benchmark NAV emails:
  1. NSE Nifty Multi Asset — ZIP attachment containing a CSV with one row per day
  2. CRISIL Composite Bond Index — Excel attachment with full historical data

Called by gmail_watcher.py when matching emails are detected.
"""

import io
import zipfile
import logging
import tempfile
import os
from datetime import datetime, date
from typing import Optional

logger = logging.getLogger(__name__)

# ── Email identifiers ─────────────────────────────────────────────────────────

NSE_SENDER          = "indices@nse.co.in"
NSE_SUBJECT_KEYWORD = "Nifty Multi Asset"

CRISIL_SENDER          = "Crisil.Indices@crisil.com"
CRISIL_SUBJECT_KEYWORD = "CRISIL Indices"
CRISIL_INDEX_NAME      = "CRISIL Composite Bond Index"


# ── NSE ZIP parser ────────────────────────────────────────────────────────────

def parse_nse_zip(file_bytes: bytes) -> Optional[tuple[date, float]]:
    """
    Parse the NSE Nifty Multi Asset ZIP email attachment.
    ZIP contains one CSV: INDEX_NAME, DATE, INDEX_VALUE
    Returns (nav_date, value) or None.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            csv_files = [f for f in z.namelist() if f.endswith(".csv")]
            if not csv_files:
                logger.warning("NSE ZIP: no CSV found inside ZIP")
                return None

            csv_bytes = z.read(csv_files[0])
            lines = csv_bytes.decode("utf-8", errors="ignore").strip().split("\n")

        if len(lines) < 2:
            logger.warning("NSE ZIP CSV: too few lines")
            return None

        # Skip header row, parse data row
        # Format: INDEX_NAME,DATE,INDEX_VALUE
        # e.g.: Nifty Multi Asset - Equity : Arbitrage : REITs/InvITs (50:40:10) Index,10-08-2026,4791.35
        data_line = lines[1].strip()
        parts = data_line.split(",")
        if len(parts) < 3:
            logger.warning(f"NSE ZIP CSV: unexpected format: {data_line}")
            return None

        # Date is second field — format DD-MM-YYYY
        date_str = parts[1].strip()
        nav_date = datetime.strptime(date_str, "%d-%m-%Y").date()

        # Value is third field
        value = float(parts[2].strip())

        logger.info(f"NSE ZIP parsed: {nav_date} = {value}")
        return nav_date, value

    except Exception as e:
        logger.error(f"NSE ZIP parse error: {e}", exc_info=True)
        return None


# ── CRISIL Excel parser ───────────────────────────────────────────────────────

def parse_crisil_excel(file_bytes: bytes) -> list[tuple[date, float]]:
    """
    Parse the CRISIL Composite Bond Index Excel attachment.
    Sheet "Index Values": Column A = Date (datetime), Column B = Index value.
    Returns list of (nav_date, value) tuples, sorted by date ascending.
    """
    try:
        import openpyxl

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        try:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        finally:
            tmp.close()

        try:
            wb = openpyxl.load_workbook(tmp_path, data_only=True, read_only=True)
        finally:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass

        # Find the "Index Values" sheet
        sheet_name = None
        for name in wb.sheetnames:
            if "index" in name.lower() and "value" in name.lower():
                sheet_name = name
                break
        if not sheet_name:
            sheet_name = wb.sheetnames[0]

        ws = wb[sheet_name]
        rows = []
        skipped_header = False

        for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
            date_val, value_val = row

            # Skip header row
            if not skipped_header:
                skipped_header = True
                if isinstance(date_val, str) and not _try_parse_date(date_val):
                    continue
                if date_val is None:
                    continue

            if date_val is None or value_val is None:
                continue

            nav_date = _try_parse_date(date_val)
            if not nav_date:
                continue

            try:
                value = float(value_val)
            except (ValueError, TypeError):
                continue

            rows.append((nav_date, value))

        rows.sort(key=lambda r: r[0])
        logger.info(f"CRISIL Excel parsed: {len(rows)} rows, "
                    f"range {rows[0][0]} to {rows[-1][0]}" if rows else "0 rows")
        return rows

    except Exception as e:
        logger.error(f"CRISIL Excel parse error: {e}", exc_info=True)
        return []


def _try_parse_date(val) -> Optional[date]:
    """Try to parse a date value from various formats."""
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    formats = ["%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y", "%d/%m/%Y", "%d-%m-%Y"]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None